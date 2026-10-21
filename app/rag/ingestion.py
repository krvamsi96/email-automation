"""
Universal Document Ingestion Pipeline for Intelligent Email Copilot.
Loads documents (PDF, DOCX, TXT, MD, CSV), cleans text, chunks them and stores
the chunks in the workspace knowledge collection, tagged with their document id.

Cleaning is format-generic: running page headers/footers (e.g. the date/URL/page
lines a browser prints on every page), PDF ligatures and line-break hyphenation
are removed, and repeated passages (quoted e-mail history, duplicated pages) are
indexed once. Each chunk is embedded together with its document title so a
passage keeps its topic even when the text itself does not repeat it.

Documents that contain e-mail correspondence are split into their individual
messages first (see rag/correspondence.py): every chunk then belongs to exactly
one message and carries its sender, date and subject, recipient address blocks
are left out, and repeated copies of a message are indexed once. Every chunk is
sized so that its embedded text fits the embedding model's input window: text
beyond that window would otherwise be silently ignored by the model.
"""

import csv
import logging
import re
import unicodedata
from bisect import bisect_right
from collections import Counter
from pathlib import Path
from typing import List, Optional, Sequence, Set

from langchain_core.documents import Document

from app.config import Config
from app.rag.correspondence import split_messages
from app.rag.vector_store import VectorStoreManager, vector_store_manager

logger = logging.getLogger(__name__)

# Extension -> display type. Legacy binary .doc is not supported by python-docx.
SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".txt": "TXT",
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".csv": "CSV",
}

PAGE_SEPARATOR = "\n\n"
MEANINGLESS_TITLES = {"", "untitled", "about:blank", "microsoft word", "document"}


class DocumentLoadError(Exception):
    """Raised when a document cannot be read. The message is safe to show to users."""


# ---------------------------------------------------------------------------
# text cleaning
# ---------------------------------------------------------------------------
def clean_text(text: str) -> str:
    """Normalises extracted text without changing its meaning."""
    text = unicodedata.normalize("NFKC", text or "")  # ligatures (ﬁ -> fi), full-width forms, odd spaces
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u200b\ufeff]", "", text)  # control chars, zero-width space, BOM
    text = re.sub(r"(?<=[a-z])-\n(?=[a-z])", "", text)  # "recommen-\ndation" -> "recommendation"
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _line_signature(line: str) -> str:
    """Line identity for header/footer detection: case-, spacing- and number-insensitive."""
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", line.strip().lower()))


def strip_running_lines(pages: Sequence[str], depth: int = 2) -> List[str]:
    """
    Removes running headers and footers. Pages are peeled one line at a time from
    the top and from the bottom (at most `depth` lines each): the outermost line is
    removed only while the same line (ignoring numbers) is outermost on most pages.
    With only two pages the evidence is weak, so just lines carrying page numbers,
    dates or times are removed there. Body text, including passages quoted on
    several pages, is kept.
    """
    if len(pages) < 2:
        return list(pages)
    split = [p.split("\n") for p in pages]
    needed = max(2, -(-3 * len(pages) // 5))  # 60% of pages, rounded up

    def qualifies(sig: str) -> bool:
        return "#" in sig or (len(pages) >= 3 and len(sig) >= 15)

    def edge_index(lines: List[str], from_top: bool) -> Optional[int]:
        order = range(len(lines)) if from_top else range(len(lines) - 1, -1, -1)
        return next((i for i in order if lines[i].strip()), None)

    for from_top in (True, False):
        for _ in range(depth):
            edge = [edge_index(lines, from_top) for lines in split]
            sigs = Counter(_line_signature(split[p][i]) for p, i in enumerate(edge) if i is not None)
            if not sigs:
                break
            sig, count = sigs.most_common(1)[0]
            if count < needed or not qualifies(sig):
                break
            for p, i in enumerate(edge):
                if i is not None and _line_signature(split[p][i]) == sig:
                    del split[p][i]
    return ["\n".join(lines).strip() for lines in split]


def _shingles(text: str, size: int = 5) -> Set[int]:
    words = re.findall(r"\w+", text.lower())
    if len(words) <= size:
        return {hash(" ".join(words))} if words else set()
    return {hash(" ".join(words[i:i + size])) for i in range(len(words) - size + 1)}


def is_near_duplicate(a: Set[int], b: Set[int], threshold: float = 0.85) -> bool:
    """True when most of the smaller passage is contained in the other one."""
    if not a or not b:
        return False
    return len(a & b) / min(len(a), len(b)) >= threshold


def document_title(meta_title: Optional[str], file_name: str) -> str:
    title = clean_text(str(meta_title or "")).replace("\n", " ").strip()
    if title.lower() in MEANINGLESS_TITLES or len(title) < 4:
        title = re.sub(r"[_\s]+", " ", Path(file_name).stem).strip()
    return title[:200]


# ---------------------------------------------------------------------------
# loaders
# ---------------------------------------------------------------------------
class UniversalDocumentLoader:
    """Loads various document formats into standardized LangChain Document objects."""

    @staticmethod
    def read_pdf(file_path: Path) -> List[Document]:
        import pypdf
        try:
            reader = pypdf.PdfReader(str(file_path))
            if reader.is_encrypted:
                raise DocumentLoadError("This PDF is password protected.")
            pages = [clean_text(page.extract_text() or "") for page in reader.pages]
            try:
                meta_title = (reader.metadata or {}).get("/Title")
            except Exception:
                meta_title = None
        except DocumentLoadError:
            raise
        except Exception as e:
            logger.error("Failed to read PDF %s: %s", file_path, e)
            raise DocumentLoadError("The PDF file could not be read. It may be damaged.") from e
        title = document_title(meta_title, file_path.name)
        docs = []
        for page_number, text in enumerate(strip_running_lines(pages), start=1):
            if text:
                docs.append(Document(
                    page_content=text,
                    metadata={"source": file_path.name, "format": "pdf", "page": page_number, "title": title},
                ))
        return docs

    @staticmethod
    def read_docx(file_path: Path) -> List[Document]:
        try:
            import docx
            doc = docx.Document(str(file_path))
        except Exception as e:
            logger.error("Failed to read DOCX %s: %s", file_path, e)
            raise DocumentLoadError("The Word document could not be read. It may be damaged.") from e
        parts = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        content = clean_text("\n".join(parts))
        if not content:
            return []
        title = document_title(getattr(doc.core_properties, "title", ""), file_path.name)
        return [Document(page_content=content, metadata={"source": file_path.name, "format": "docx", "title": title})]

    @staticmethod
    def read_text(file_path: Path, fmt: str) -> List[Document]:
        try:
            text = file_path.read_text(encoding="utf-8-sig", errors="ignore")
        except OSError as e:
            raise DocumentLoadError("The file could not be opened.") from e
        text = clean_text(text)
        if not text:
            return []
        return [Document(page_content=text, metadata={"source": file_path.name, "format": fmt,
                                                      "title": document_title("", file_path.name)})]

    @staticmethod
    def read_csv(file_path: Path) -> List[Document]:
        docs = []
        title = document_title("", file_path.name)
        try:
            with open(file_path, mode="r", encoding="utf-8-sig", errors="ignore", newline="") as f:
                reader = csv.DictReader(f)
                for i, row in enumerate(reader):
                    row_str = " | ".join(f"{k}: {v}" for k, v in row.items() if k and v)
                    if row_str.strip():
                        docs.append(Document(
                            page_content=clean_text(row_str),
                            metadata={"source": file_path.name, "format": "csv", "row": i, "title": title},
                        ))
        except (OSError, csv.Error) as e:
            logger.error("Failed to read CSV %s: %s", file_path, e)
            raise DocumentLoadError("The CSV file could not be read.") from e
        return docs

    @classmethod
    def load_file_strict(cls, file_path: Path) -> List[Document]:
        """Loads a file, raising DocumentLoadError with a user-friendly message on failure."""
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return cls.read_pdf(file_path)
        if suffix == ".docx":
            return cls.read_docx(file_path)
        if suffix == ".txt":
            return cls.read_text(file_path, "txt")
        if suffix in (".md", ".markdown"):
            return cls.read_text(file_path, "md")
        if suffix == ".csv":
            return cls.read_csv(file_path)
        raise DocumentLoadError(f"Files of type '{suffix or 'unknown'}' are not supported.")

    # Lenient loaders kept for backwards compatibility: return [] on failure.
    @classmethod
    def _lenient(cls, fn, *args) -> List[Document]:
        try:
            return fn(*args)
        except DocumentLoadError as e:
            logger.error("%s", e)
            return []

    @classmethod
    def load_pdf(cls, file_path: Path) -> List[Document]:
        return cls._lenient(cls.read_pdf, file_path)

    @classmethod
    def load_docx(cls, file_path: Path) -> List[Document]:
        return cls._lenient(cls.read_docx, file_path)

    @classmethod
    def load_txt(cls, file_path: Path) -> List[Document]:
        return cls._lenient(cls.read_text, file_path, "txt")

    @classmethod
    def load_md(cls, file_path: Path) -> List[Document]:
        return cls._lenient(cls.read_text, file_path, "md")

    @classmethod
    def load_csv(cls, file_path: Path) -> List[Document]:
        return cls._lenient(cls.read_csv, file_path)

    @classmethod
    def load_file(cls, file_path: Path) -> List[Document]:
        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            logger.info("Skipping unsupported file format: %s", file_path.name)
            return []
        return cls._lenient(cls.load_file_strict, file_path)


# ---------------------------------------------------------------------------
# pipeline
# ---------------------------------------------------------------------------
def _topic_key(text: str) -> str:
    """Subject/title comparison key: no Re:/Fwd: prefixes, punctuation or case."""
    text = re.sub(r"^\W*((re|fwd?|fw)\b\W*)+", "", (text or "").lower())
    return " ".join(re.findall(r"\w+", text))


def chunk_heading(metadata: dict) -> str:
    """Context a chunk is embedded and keyword-indexed with: document title, message subject and sender."""
    title = metadata.get("title") or ""
    subject = metadata.get("subject") or ""
    lines = [title] if title else []
    if subject and _topic_key(subject) not in _topic_key(title):
        lines.append(subject)
    if metadata.get("sender"):
        lines.append(f"From: {metadata['sender']}")
    return "\n".join(lines)


def embedding_text(chunk: Document) -> str:
    """What is embedded for a chunk: its heading (title, subject, sender) followed by its text."""
    heading = chunk_heading(chunk.metadata)
    return f"{heading}\n\n{chunk.page_content}" if heading else chunk.page_content


class IngestionPipeline:
    """Coordinates document loading, splitting, and vector store population."""

    def __init__(
        self,
        knowledge_base_path: Optional[str] = None,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
        manager: Optional[VectorStoreManager] = None,
    ):
        self.knowledge_base_path = Path(knowledge_base_path or Config.KNOWLEDGE_BASE_PATH)
        self.chunk_size = chunk_size or Config.CHUNK_SIZE
        self.chunk_overlap = Config.CHUNK_OVERLAP if chunk_overlap is None else chunk_overlap
        self.manager = manager or vector_store_manager
        self._splitters = None

    def _get_splitters(self):
        # Imported on first use: the text-splitters package pulls in heavy ML libraries at import time.
        if self._splitters is None:
            from langchain_text_splitters import Language, RecursiveCharacterTextSplitter

            self._splitters = (
                RecursiveCharacterTextSplitter(
                    chunk_size=self.chunk_size, chunk_overlap=self.chunk_overlap,
                    separators=["\n\n", "\n", ". ", " ", ""], add_start_index=True,
                ),
                RecursiveCharacterTextSplitter.from_language(
                    Language.MARKDOWN, chunk_size=self.chunk_size, chunk_overlap=self.chunk_overlap, add_start_index=True,
                ),
            )
        return self._splitters

    @property
    def text_splitter(self):
        return self._get_splitters()[0]

    @property
    def markdown_splitter(self):
        return self._get_splitters()[1]

    def _split(self, raw_docs: List[Document]) -> List[Document]:
        """
        Pages of one document are joined before splitting so a passage that runs
        over a page break stays together; each chunk keeps the page it starts on.
        Row-based sources (CSV) keep one record per chunk. Correspondence is split
        per message (see _split_messages).
        """
        if not raw_docs:
            return []
        if raw_docs[0].metadata.get("format") == "csv":
            return self.text_splitter.split_documents(raw_docs)
        offsets: List[int] = []
        pages: List[Optional[int]] = []
        parts: List[str] = []
        position = 0
        for d in raw_docs:
            offsets.append(position)
            pages.append(d.metadata.get("page"))
            parts.append(d.page_content)
            position += len(d.page_content) + len(PAGE_SEPARATOR)
        meta = {k: v for k, v in raw_docs[0].metadata.items() if k != "page"}
        text = PAGE_SEPARATOR.join(parts)

        def page_at(offset: int) -> Optional[int]:
            return pages[max(0, bisect_right(offsets, offset) - 1)]

        chunks = self._split_messages(text, meta, page_at)
        if chunks is not None:
            return chunks
        splitter = self.markdown_splitter if meta.get("format") == "md" else self.text_splitter
        chunks = splitter.split_documents([Document(page_content=text, metadata=meta)])
        for chunk in chunks:
            page = page_at(chunk.metadata.get("start_index", 0))
            if page is not None:
                chunk.metadata["page"] = page
        return chunks

    def _split_messages(self, text: str, meta: dict, page_at) -> Optional[List[Document]]:
        """
        Chunks for a document that contains e-mail messages, or None when it does not.
        Each message is split on its own and its chunks carry sender, date, subject,
        a sortable timestamp and the message number; repeated copies are skipped.
        """
        messages = split_messages(text, meta.get("title", ""))
        if not messages:
            return None
        chunks: List[Document] = []
        seen: List[Set[int]] = []
        for msg in messages:
            body = msg.body
            sig = _shingles(body)
            if any(is_near_duplicate(sig, s) for s in seen):
                continue  # quoted or re-printed copy of a message already indexed
            seen.append(sig)
            info = {"sender": msg.sender, "date": msg.date, "subject": msg.subject,
                    "timestamp": msg.timestamp, "message": len(seen) - 1}
            for chunk in self.text_splitter.split_documents([Document(page_content=body, metadata=dict(meta))]):
                chunk.metadata.update({k: v for k, v in info.items() if v != ""})
                page = page_at(msg.offset + chunk.metadata.get("start_index", 0))
                if page is not None:
                    chunk.metadata["page"] = page
                chunks.append(chunk)
        return chunks

    def _fit_to_window(self, chunks: List[Document]) -> List[Document]:
        """
        Re-splits chunks whose embedded text (heading + text) is longer than the
        embedding model's input window, so no part of any chunk goes unembedded.
        """
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        from app.rag.embeddings import get_embeddings

        embeddings = get_embeddings()
        limit = embeddings.max_tokens - 2
        fitted: List[Document] = []
        pending = list(chunks)
        while pending:
            chunk = pending.pop(0)
            tokens = embeddings.count_tokens(embedding_text(chunk))
            if tokens <= limit or len(chunk.page_content) < 200:
                fitted.append(chunk)
                continue
            heading = embeddings.count_tokens(chunk_heading(chunk.metadata))
            share = max(0.2, (limit - heading) / max(1, tokens - heading))
            size = max(150, int(len(chunk.page_content) * share * 0.9))
            splitter = RecursiveCharacterTextSplitter(
                chunk_size=size, chunk_overlap=min(self.chunk_overlap, size // 5),
                separators=["\n\n", "\n", ". ", " ", ""])
            parts = splitter.split_text(chunk.page_content)
            if len(parts) < 2:
                fitted.append(chunk)
                continue
            pending[:0] = [Document(page_content=p, metadata=dict(chunk.metadata)) for p in parts]
        return fitted

    def chunk_file(self, file_path: Path, doc_id: str, display_name: Optional[str] = None) -> List[Document]:
        """Loads, cleans and splits one file; every chunk is tagged with doc_id. Repeated passages are dropped."""
        raw_docs = UniversalDocumentLoader.load_file_strict(file_path)
        chunks: List[Document] = []
        seen: List[Set[int]] = []
        for chunk in self._fit_to_window(self._split(raw_docs)):
            text = chunk.page_content.strip()
            if len(re.findall(r"\w+", text)) < 3:
                continue  # stray fragments (lone page numbers, separators)
            sig = _shingles(text)
            if any(is_near_duplicate(sig, s) for s in seen):
                continue
            seen.append(sig)
            chunk.page_content = text
            chunk.metadata.pop("start_index", None)
            chunk.metadata["doc_id"] = doc_id
            chunk.metadata["source"] = display_name or file_path.name
            chunk.metadata["chunk"] = len(chunks)
            chunks.append(chunk)
        return chunks

    def store_chunks(self, doc_id: str, chunks: List[Document], manager: Optional[VectorStoreManager] = None) -> int:
        """Replaces a document's indexed chunks with the given ones."""
        manager = manager or self.manager
        from app.rag.embeddings import get_embeddings

        vectors = get_embeddings().embed_documents([embedding_text(c) for c in chunks])
        metadatas = [{k: v for k, v in c.metadata.items() if v is not None} for c in chunks]
        manager.delete_document(doc_id)
        return manager.add_chunks([f"{doc_id}-{i}" for i in range(len(chunks))],
                                  [c.page_content for c in chunks], metadatas, vectors)

    def index_file(self, file_path: Path, doc_id: str, display_name: Optional[str] = None) -> int:
        """
        (Re)indexes a single document: removes its old chunks, then stores the new ones.
        Raises DocumentLoadError when the file has no readable text.
        """
        chunks = self.chunk_file(file_path, doc_id, display_name)
        if not chunks:
            raise DocumentLoadError(
                "No readable text was found in this document. Scanned images are not supported."
            )
        return self.store_chunks(doc_id, chunks)

    def load_documents_from_directory(self) -> List[Document]:
        """Loads all supported documents from the configured knowledge directory."""
        if not self.knowledge_base_path.exists():
            logger.warning("Knowledge base directory does not exist: %s", self.knowledge_base_path)
            return []
        all_docs = []
        for file_path in sorted(self.knowledge_base_path.glob("*")):
            if file_path.is_file():
                all_docs.extend(UniversalDocumentLoader.load_file(file_path))
        logger.info("Loaded %d documents from %s", len(all_docs), self.knowledge_base_path)
        return all_docs

    def run(self, reset_existing: bool = False) -> int:
        """
        Indexes every document in the default workspace's knowledge folder.
        Documents are registered in the Knowledge Base so they appear in the app
        and can later be deleted individually. Returns the number of chunks stored.
        """
        from app.services.knowledge import knowledge_service
        from app.rag.vector_store import DEFAULT_WORKSPACE_ID

        return knowledge_service.reindex_all(DEFAULT_WORKSPACE_ID, reset=reset_existing)


# Global singleton ingestion pipeline
ingestion_pipeline = IngestionPipeline()
